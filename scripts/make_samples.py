"""Generate the sample .docx fixtures used by parser verification.

Run: .venv/bin/python scripts/make_samples.py
"""

from pathlib import Path

import docx
from docx.oxml import parse_xml

SAMPLES = Path(__file__).resolve().parents[1] / "samples"

_OMML = (
    '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
    "<m:r><m:t>{expr}</m:t></m:r></m:oMath>"
)


def make_plain() -> None:
    d = docx.Document()
    d.add_paragraph("Hello handwriting world! This is sample 1, typed on 2026-07-12.")
    d.save(SAMPLES / "plain.docx")


def make_multi() -> None:
    d = docx.Document()
    d.add_paragraph("First paragraph with numbers 123 and (parentheses).")
    p = d.add_paragraph("Second paragraph, first line.")
    run = p.add_run()
    run.add_break()
    run.add_text("Second paragraph, after a soft line break.")
    d.add_paragraph("")  # blank line between paragraphs
    d.add_paragraph(
        "Word-style “smart quotes”, an em—dash, "
        "a tab\there, and café."
    )
    d.save(SAMPLES / "multi.docx")


def make_equation() -> None:
    d = docx.Document()
    d.add_paragraph("Text before the equation.")
    p = d.add_paragraph("Inline equation here: ")
    p._p.append(parse_xml(_OMML.format(expr="E=mc^2")))
    display = d.add_paragraph("")  # display equation on its own paragraph
    display._p.append(parse_xml(_OMML.format(expr="x = (-b + sqrt(b^2-4ac)) / 2a")))
    d.add_paragraph("Text after the equation.")
    d.save(SAMPLES / "equation.docx")


def main() -> None:
    SAMPLES.mkdir(exist_ok=True)
    make_plain()
    make_multi()
    make_equation()
    print(f"wrote plain.docx, multi.docx, equation.docx into {SAMPLES}")


if __name__ == "__main__":
    main()
