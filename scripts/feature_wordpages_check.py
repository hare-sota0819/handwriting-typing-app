"""Feature verification: mirror Word's own pages (style.word_pages).

Word records where its last layout began each page (w:lastRenderedPageBreak).
With word_pages on, what sits on one Word page starts on a fresh handwriting
page, so blank lines typed to push content down need no re-fitting.

Checks:
- Parser: a marker at the start of a paragraph flags it (word_page_start); a
  marker inside a paragraph becomes a WordPageBreak run; the text is unchanged.
- OFF (the default): a marked document lays out exactly like its unmarked
  twin — the markers are inert.
- ON, marker between paragraphs: the marked paragraph opens the next page at
  the top; blank paragraphs padding the end of the Word page are dropped even
  when there are more of them than a page holds.
- ON, marker inside a paragraph: the line being written is finished, then the
  page ends; the rest of the paragraph continues on the next page.
- ON, a Word page too long for one handwriting page spills over, and the next
  Word page still opens a fresh page.
- A marker right after an explicit page break opens one new page, not two.
Exit 0 = pass.
"""

import sys
from dataclasses import replace
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import docx
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement

from handwriting.document_model import TextRun, WordPageBreak
from handwriting.glyphstore import GlyphStore
from handwriting.layout import PageSpec, layout_document
from handwriting.parser import parse_docx
from handwriting.style import HandwritingStyle
from handwriting.synthglyphs import DEFAULT_STORE

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "out" / "wordpages"

LONG = " ".join(f"word{i}" for i in range(12))  # a few handwriting lines


def _marker():
    return OxmlElement("w:lastRenderedPageBreak")


def _add_marked(paragraph, text: str) -> None:
    """A run that starts a Word page: the marker, then its text."""
    run = paragraph.add_run()
    run._r.append(_marker())
    t = OxmlElement("w:t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    run._r.append(t)


_M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _add_marked_equation(paragraph, display: bool) -> None:
    """An equation whose first run carries the marker — where Word puts it
    when a page begins with an equation (seen in the user's prelab report)."""
    from lxml import etree
    omath = (f'<m:oMath xmlns:m="{_M_NS}" xmlns:w="{_W_NS}"><m:r>'
             '<w:lastRenderedPageBreak/><m:t>x</m:t></m:r></m:oMath>')
    if display:
        omath = (f'<m:oMathPara xmlns:m="{_M_NS}" xmlns:w="{_W_NS}">'
                 f'{omath}</m:oMathPara>')
    paragraph._p.append(etree.fromstring(omath))


def build(name: str, spec: list, marked: bool) -> Path:
    """spec items: str = paragraph, None = blank paragraph, ("mark", str) =
    paragraph opening a Word page, ("mid", a, b) = paragraph whose Word page
    ends after a, ("break",) = explicit page break."""
    d = docx.Document()
    for item in spec:
        if item is None:
            d.add_paragraph("")
        elif isinstance(item, str):
            d.add_paragraph(item)
        elif item[0] == "mark":
            p = d.add_paragraph()
            if marked:
                _add_marked(p, item[1])
            else:
                p.add_run(item[1])
        elif item[0] == "mid":
            p = d.add_paragraph()
            p.add_run(item[1] + " ")
            if marked:
                _add_marked(p, item[2])
            else:
                p.add_run(item[2])
        elif item[0] == "break":
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    path = OUT / f"{name}{'' if marked else '_plain'}.docx"
    d.save(str(path))
    return path


def lines_of(page, style) -> list[list]:
    """A page's glyphs grouped into written lines (by baseline proximity)."""
    rows: list[list] = []
    for g in page.glyphs:
        if rows and abs(rows[-1][-1].y - g.y) < 0.5 * style.base_size_px:
            rows[-1].append(g)
        else:
            rows.append([g])
    return rows


def text_of(page) -> str:
    return "".join(g.char for g in page.glyphs)


def main() -> int:
    failures: list[str] = []
    OUT.mkdir(parents=True, exist_ok=True)
    store = GlyphStore(DEFAULT_STORE)
    off = HandwritingStyle()
    on = replace(off, word_pages=True)
    spec = PageSpec()
    top = spec.margin_px + 2 * off.base_size_px

    def layout(path, style):
        return layout_document(parse_docx(path), store, style, spec, seed=2)

    # --- parser
    between = [LONG, None, None, ("mark", "Second page starts here."), LONG]
    doc = parse_docx(build("between", between, True))
    flags = [p.word_page_start for p in doc.paragraphs]
    if flags != [False, False, False, True, False]:
        failures.append(f"parser: word_page_start flags {flags}")
    if any(isinstance(r, WordPageBreak) for p in doc.paragraphs for r in p.runs):
        failures.append("parser: a paragraph-start marker leaked into the runs")

    mid = [("mid", "alpha beta gamma", LONG)]
    doc = parse_docx(build("mid", mid, True))
    kinds = [type(r).__name__ for r in doc.paragraphs[0].runs]
    if kinds != ["TextRun", "WordPageBreak", "TextRun"]:
        failures.append(f"parser: mid-paragraph marker gave runs {kinds}")
    plain = parse_docx(build("mid", mid, False))
    joined = " ".join(r.text for r in doc.paragraphs[0].runs if isinstance(r, TextRun))
    if joined != plain.paragraphs[0].runs[0].text:
        failures.append("parser: a marker changed the paragraph's text")

    # A Word page that begins with an equation carries its marker INSIDE the
    # equation. Missed, two Word pages of the user's report merged into one.
    d = docx.Document()
    d.add_paragraph("First Word page.")
    _add_marked_equation(d.add_paragraph(), display=True)
    inline = d.add_paragraph()
    inline.add_run("so that ")
    _add_marked_equation(inline, display=False)
    path = OUT / "in_equation.docx"
    d.save(str(path))
    doc = parse_docx(path)
    shape = [(p.word_page_start, [type(r).__name__ for r in p.runs])
             for p in doc.paragraphs]
    want = [(False, ["TextRun"]), (True, ["MathBlock"]),
            (False, ["TextRun", "WordPageBreak", "MathBlock"])]
    if shape != want:
        failures.append(f"parser: marker inside an equation gave {shape}")

    # --- OFF: markers are inert
    for name, s in (("between", between), ("mid", mid)):
        a = layout(build(name, s, True), off)
        b = layout(build(name, s, False), off)
        if [p.glyphs for p in a] != [p.glyphs for p in b]:
            failures.append(f"off: marked '{name}' lays out differently from its twin")

    # --- ON: marker between paragraphs
    pages = layout(build("between", between, True), on)
    if len(pages) != 2:
        failures.append(f"between: want 2 pages, got {len(pages)}")
    else:
        if "Secondpage" in text_of(pages[0]):
            failures.append("between: the marked paragraph stayed on page one")
        if not text_of(pages[1]).startswith("Secondpage"):
            failures.append(f"between: page two starts {text_of(pages[1])[:12]!r}")
        if pages[1].glyphs[0].y > top:
            failures.append("between: page two does not start at the top")

    # padding blanks: more than a page holds must not cost a page
    padded = [LONG] + [None] * 80 + [("mark", "Second page starts here.")]
    pages = layout(build("padded", padded, True), on)
    if len(pages) != 2:
        failures.append(f"padded: want 2 pages, got {len(pages)}")
    elif pages[1].glyphs[0].y > top:
        failures.append("padded: page two does not start at the top")

    # --- ON: marker inside a paragraph
    pages = layout(build("mid", mid, True), on)
    if len(pages) != 2:
        failures.append(f"mid: want 2 pages, got {len(pages)}")
    else:
        rows = lines_of(pages[0], on)
        first = text_of(pages[0])
        if len(rows) != 1:
            failures.append(f"mid: page one holds {len(rows)} lines, want the one "
                            "line being written")
        if not first.startswith("alphabetagamma") or len(first) <= len("alphabetagamma"):
            failures.append(f"mid: the line was not finished before the page "
                            f"ended: {first!r}")
        whole = text_of(layout(build("mid", mid, False), off)[0])
        if first + text_of(pages[1]) != whole:
            failures.append("mid: text lost or reordered across the page end")
        if pages[1].glyphs[0].y > top:
            failures.append("mid: page two does not start at the top")

    # --- ON: a Word page longer than a handwriting page spills over
    huge = " ".join(f"word{i}" for i in range(150))
    spill = [huge, ("mark", "Next Word page.")]
    pages = layout(build("spill", spill, True), on)
    n_flow = len(layout(build("spill", [huge], False), off))
    if n_flow < 2:
        failures.append("spill: fixture too short to overflow a page")
    if len(pages) != n_flow + 1:
        failures.append(f"spill: want {n_flow + 1} pages, got {len(pages)}")
    elif not text_of(pages[-1]).startswith("NextWord") or pages[-1].glyphs[0].y > top:
        failures.append("spill: the next Word page did not open a fresh page")

    # --- explicit break followed by a marker: one new page
    both = ["First page.", ("break",), ("mark", "After the break.")]
    pages = layout(build("both", both, True), on)
    if len(pages) != 2:
        failures.append(f"explicit+marker: want 2 pages, got {len(pages)}")

    if failures:
        print(f"FAIL ({len(failures)} problems):")
        for f in failures:
            print(f"  - {f}")
        return 1
    print("PASS: Word-page mirroring checks all green")
    return 0


if __name__ == "__main__":
    sys.exit(main())
