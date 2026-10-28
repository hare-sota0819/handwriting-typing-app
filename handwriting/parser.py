""".docx → Document Model parser: text runs and equations, in order.

Paragraphs without equations take python-docx's paragraph text (the Phase 1
path, unchanged). A paragraph with OMML in it is walked child by child so
text and MathBlock nodes interleave exactly as Word shows them; the LaTeX
for each block is attached later by mathrender.py, paired by document-order
index.
"""

import logging
import re
import string
from pathlib import Path

import docx

from .document_model import (Document, MathBlock, Paragraph, TextRun,
                             WordPageBreak)

logger = logging.getLogger(__name__)

_OMML_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# Word autocorrect / typographic characters mapped onto the glyph charset.
_CHAR_NORMALIZE = {
    ord("‘"): "'", ord("’"): "'",   # smart single quotes
    ord("“"): '"', ord("”"): '"',   # smart double quotes
    ord("–"): "-", ord("—"): "-",   # en/em dash
    ord("…"): "...",                     # ellipsis
    ord(" "): " ",                       # non-breaking space
    ord("\t"): " ",
}
# Enclosed digits (Word's ① ② … list markers): no glyphs exist for them and
# they were silently dropped, so write them the way a hand does — "(1)".
_CHAR_NORMALIZE.update({0x2460 + i: f"({i + 1})" for i in range(20)})   # ①-⑳
_CHAR_NORMALIZE.update({0x2474 + i: f"({i + 1})" for i in range(20)})   # ⑴-⒇
_CHAR_NORMALIZE.update({0x2488 + i: f"{i + 1}." for i in range(20)})    # ⒈-⒛
_CHAR_NORMALIZE[0x24EA] = "(0)"                                          # ⓪
# Greek "symbol" codepoints typed in prose (ϵ, ϕ, ...): the charset holds the
# ordinary letter, so map them the same way the equation renderer aliases them.
_CHAR_NORMALIZE.update({ord(v): g for v, g in
                        (("ϵ", "ε"), ("ϕ", "φ"), ("ϑ", "θ"), ("ϱ", "ρ"),
                         ("ϖ", "π"), ("ϰ", "κ"), ("ς", "σ"))})
# Relations with no glyph of their own, written as the nearest one collected
# (a ≳ in the user's ch13 was spaced over and silently lost).
_CHAR_NORMALIZE.update({ord(v): g for v, g in
                        (("≳", "≥"), ("⩾", "≥"), ("≧", "≥"),
                         ("≲", "≤"), ("⩽", "≤"), ("≦", "≤"))})

# Characters the pipeline is built for; anything else is kept but logged so
# the glyph-selection stage can apply its fallback knowingly. Unicode
# superscripts are handled by layout (drawn as raised, shrunken digits).
_KNOWN_CHARS = set(
    string.ascii_letters + string.digits + string.punctuation + " "
    + "⁰¹²³⁴⁵⁶⁷⁸⁹⁻" + "·×÷"
)


def _normalize_text(text: str) -> str:
    text = text.translate(_CHAR_NORMALIZE)
    text = re.sub(r" {2,}", " ", text)
    return text.strip()


_INDEX_ATTR = "{urn:handwriting}index"


def parse_equations(body) -> int:
    """Number every m:oMath in the body, in document order.

    Stamps each element with its index (an in-memory attribute; the file is
    never written back). pandoc extracts LaTeX in this same order, so the
    index is what pairs a MathBlock with its LaTeX (mathrender.py).
    Equations inside tables/text boxes count too, even though only body
    paragraphs are laid out — skipping them would shift the pairing.
    """
    found = body.findall(f".//{_OMML_NS}oMath")
    if found:
        logger.info("%d equation(s) found", len(found))
    for i, el in enumerate(found):
        el.set(_INDEX_ATTR, str(i))
    return len(found)


# Word's equation justifications: "centerGroup" centres the equations of a
# math paragraph as one block, which for our one-equation lines is centring.
_JC_ALIGN = {"left": "left", "right": "right",
             "center": "center", "centerGroup": "center"}


def _jc_align(jc, default: str) -> str:
    """Our alignment for an m:jc / m:defJc element (None -> default)."""
    if jc is None:
        return default
    return _JC_ALIGN.get(jc.get(f"{_OMML_NS}val"), default)


def default_math_align(docx_doc) -> str:
    """The document-wide display equation justification (settings m:defJc)."""
    jc = docx_doc.settings.element.find(f"{_OMML_NS}mathPr/{_OMML_NS}defJc")
    return _jc_align(jc, "center")


def _walk_paragraph(el, in_display: bool, align: str):
    """Yield ("text", str) / ("math", MathBlock) for a paragraph's content."""
    for child in el.iterchildren():
        tag = child.tag
        if tag == f"{_OMML_NS}oMathPara":
            jc = child.find(f"{_OMML_NS}oMathParaPr/{_OMML_NS}jc")
            yield from _walk_paragraph(child, True, _jc_align(jc, align))
        elif tag == f"{_OMML_NS}oMath":
            from lxml import etree
            # A Word page that begins with an equation keeps its marker inside
            # the equation's first run; the page begins before the equation.
            if child.find(f".//{_W_NS}lastRenderedPageBreak") is not None:
                yield ("pagehint", None)
            yield ("math", MathBlock(
                omml_xml=etree.tostring(child, encoding="unicode"),
                display=in_display, index=int(child.get(_INDEX_ATTR)),
                align=align))
        elif tag == f"{_W_NS}t":
            yield ("text", child.text or "")
        elif tag == f"{_W_NS}tab":
            yield ("text", "\t")
        elif tag == f"{_W_NS}noBreakHyphen":
            yield ("text", "-")
        elif tag == f"{_W_NS}br" and child.get(f"{_W_NS}type") == "page":
            yield ("pagebreak", None)
        elif tag == f"{_W_NS}lastRenderedPageBreak":
            yield ("pagehint", None)
        elif tag in (f"{_W_NS}br", f"{_W_NS}cr"):
            yield ("text", "\n")
        elif tag in (f"{_W_NS}pPr", f"{_W_NS}rPr", f"{_W_NS}del",
                     f"{_OMML_NS}oMathParaPr"):
            continue
        else:  # w:r, w:hyperlink, w:ins, w:sdt, w:smartTag, ...: descend
            yield from _walk_paragraph(child, in_display, align)


def _math_paragraph(para_el, unknown: set[str],
                    math_align: str = "center") -> list[Paragraph]:
    """Paragraphs (split at soft breaks) for a paragraph holding equations.

    math_align is the document's default display justification, used where
    an equation's own m:oMathPara sets none.
    """
    pieces = list(_walk_paragraph(para_el, False, math_align))
    has_text = any(k == "text" and t.strip() for k, t in pieces)
    out: list[Paragraph] = []
    runs: list[TextRun | MathBlock] = []
    buf = ""
    break_next = _page_break_before(para_el)
    hint_next = False  # Word's last layout began a page at the next paragraph

    def flush_text() -> None:
        nonlocal buf
        text = _normalize_text(buf)
        buf = ""
        if text:
            unknown.update(c for c in text if c not in _KNOWN_CHARS)
            runs.append(TextRun(text=text))

    def close_paragraph() -> None:
        nonlocal runs, break_next, hint_next, open_math
        flush_text()
        out.append(Paragraph(runs=runs, page_break_before=break_next,
                             word_page_start=hint_next))
        runs = []
        break_next = False
        hint_next = False
        open_math = None

    # The equation just placed, until the text that follows it says whether
    # Word left a space after it (_normalize_text strips, so the answer has to
    # be read off the raw run text).
    open_math: MathBlock | None = None

    for kind, item in pieces:
        if kind == "pagebreak":
            # Text before the break stays on this page; what follows starts
            # a new paragraph flagged to open a new page.
            if runs or buf.strip():
                close_paragraph()
            break_next = True
            continue
        if kind == "pagehint":
            # Word began a page here. Ahead of any content it marks the
            # paragraph; inside one it keeps its place among the runs.
            if runs or buf.strip():
                flush_text()
                runs.append(WordPageBreak())
            else:
                hint_next = True
            continue
        if kind == "math":
            item.tight_left = bool(buf) and not buf[-1].isspace()
            flush_text()
            # A lone equation in a paragraph of its own is display maths even
            # when Word did not wrap it in oMathPara.
            if not has_text:
                item.display = True
            runs.append(item)
            open_math = item
            continue
        for j, line in enumerate(item.split("\n")):
            if j > 0:
                close_paragraph()
            if line:
                if open_math is not None:
                    open_math.tight_right = not line[0].isspace()
                    open_math = None
                buf += line
    close_paragraph()
    return out


def _page_break_before(para_el) -> bool:
    return para_el.find(f"{_W_NS}pPr/{_W_NS}pageBreakBefore") is not None


def _has_page_break(para_el) -> bool:
    if _page_break_before(para_el):
        return True
    for br in para_el.iter(f"{_W_NS}br"):
        if br.get(f"{_W_NS}type") == "page":
            return True
    return False


def parse_docx(path: str | Path) -> Document:
    """Parse a plain-text .docx into the Document Model.

    - one Word paragraph → one Paragraph (empty ones preserved as blank lines)
    - soft line breaks (w:br, surfaced as '\\n') split into separate Paragraphs
    - whitespace normalized: tabs/nbsp → space, runs of spaces collapsed, ends
      stripped; smart quotes/dashes/ellipsis mapped to ASCII
    - unsupported characters are kept and logged, never a crash
    """
    docx_doc = docx.Document(str(path))
    parse_equations(docx_doc.element.body)
    math_align = default_math_align(docx_doc)

    document = Document()
    unknown: set[str] = set()
    for para in docx_doc.paragraphs:
        if (para._p.find(f".//{_OMML_NS}oMath") is not None
                or _has_page_break(para._p)
                or para._p.find(f".//{_W_NS}lastRenderedPageBreak") is not None):
            document.paragraphs.extend(
                _math_paragraph(para._p, unknown, math_align))
            continue
        for line in para.text.split("\n"):
            text = _normalize_text(line)
            unknown.update(c for c in text if c not in _KNOWN_CHARS)
            runs = [TextRun(text=text)] if text else []
            document.paragraphs.append(Paragraph(runs=runs))
    if unknown:
        logger.warning(
            "unsupported characters (kept; glyph fallback decides later): %s",
            " ".join(sorted(unknown)),
        )
    return document
